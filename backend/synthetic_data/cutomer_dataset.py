from openpyxl import Workbook
import random
from datetime import datetime, timedelta

wb = Workbook()

# Create sheets
customers_ws = wb.active
customers_ws.title = "Customers"
accounts_ws = wb.create_sheet("Accounts")
transactions_ws = wb.create_sheet("Transactions")
alerts_ws = wb.create_sheet("Alerts")
kyc_ws = wb.create_sheet("KYC")

# Headers
customers_ws.append([
    "Customer_ID", "Name", "Age", "Country", "Occupation",
    "Annual_Income", "KYC_Level", "PEP_Flag", "AML_Risk",
    "Account_Open_Date", "Customer_Since"
])

accounts_ws.append([
    "Account_ID", "Customer_ID", "Account_Type", "Currency",
    "Branch", "Balance", "Account_Open_Date", "Status"
])

transactions_ws.append([
    "Transaction_ID", "Account_ID", "Customer_ID", "Date",
    "Amount", "Currency", "Transaction_Type", "Counterparty",
    "Channel", "Location"
])

alerts_ws.append([
    "Alert_ID", "Customer_ID", "Account_ID", "Alert_Date",
    "Alert_Type", "Severity", "Reason"
])

kyc_ws.append([
    "Customer_ID", "KYC_Review_Date", "KYC_Status",
    "Documents_Verified", "Risk_Rating", "Review_Notes"
])

# Reference data
names = ["Aarav", "Vihaan", "Ananya", "Diya", "Ovian", "Riya","Elon", "Amit", "Nehal", "Priya", "Sara", "Karan",
         "Sneha", "Ishita", "Kabir", "Aang", "Sokka", "Zuko", "Azula", "Toph", "Katara", "Ty Lee", "Ozai",
         "Da Lee", "Nancy", "Iroh", "Kyoshi", "Bomi", "Weylin", "Gyatso", "Danny", "Kabir", "Sam", "Anna", "Hakoda"]
occupations = ["Engineer", "Doctor", "Teacher", "Consultant", "Trader", "Business", "Student", "Lawyer", "Analyst"]
account_types = ["Savings", "Current", "Salary", "NRI", "Loan"]
transaction_types = ["Credit", "Debit", "Transfer", "Payment"]
channels = ["Online", "Branch", "ATM", "Mobile"]
kyc_levels = ["Basic", "Standard", "Enhanced"]
risk_levels = ["Low", "Medium", "High"]

country_currency = {
    "India": "INR",
    "UAE": "AED",
    "Singapore": "SGD",
    "UK": "GBP",
    "USA": "USD",
    "Poland": "PLN",
    "China": "CNY",
    "Pakistan": "PKR",
    "Iran": "IRR",
    "Iraq": "IQD",
    "Russia": "RUB",
    "Afghanistan": "AFN",
    "Ukraine": "UAH",
    "Bahrain": "BHD",
    "Brazil": "BRL"
}

country_branches = {
    "India": ["Mumbai", "Delhi", "Bangalore", "Chennai", "Kolkata"],
    "UAE": ["Dubai", "Abu Dhabi", "Sharjah"],
    "Singapore": ["Singapore", "Jurong"],
    "UK": ["London", "Manchester", "Birmingham"],
    "USA": ["New York", "Los Angeles", "Chicago"],
    "Poland": ["Warsaw", "Krakow", "Gdansk"],
    "China": ["Beijing", "Shanghai", "Shenzhen"],
    "Pakistan": ["Karachi", "Lahore", "Islamabad"],
    "Iran": ["Tehran", "Isfahan"],
    "Iraq": ["Baghdad", "Basra"],
    "Russia": ["Moscow", "St. Petersburg"],
    "Afghanistan": ["Kabul"],
    "Ukraine": ["Kyiv", "Kharkiv"],
    "Bahrain": ["Manama"],
    "Brazil": ["Sao Paulo", "Rio de Janeiro", "Brasilia"]
}

countries = list(country_currency.keys())

start_date = datetime(2015, 1, 1)
customer_count = 8000

customer_ids = []
transaction_rows = []
alert_rows = []
kyc_rows = []

account_counter = 1
transaction_counter = 1
alert_counter = 1

for i in range(1, customer_count + 1):
    customer_id = f"CUST{i:06d}"
    customer_ids.append(customer_id)
    open_date = start_date + timedelta(days=random.randint(0, 4000))
    pep_flag = random.random() < 0.03
    esg_flag = random.random() < 0.05
    aml_risk = random.choice(["Low", "Medium", "High"] if pep_flag or esg_flag else ["Low", "Medium"])

    customer_country = random.choice(countries)
    customer_currency = country_currency[customer_country]
    customer_branches = country_branches[customer_country]

    customers_ws.append([
        customer_id,
        f"{random.choice(names)} {random.choice(['S', 'G', 'P', 'T', 'K', '', 'R'])}",
        random.randint(21, 75),
        customer_country,
        random.choice(occupations),
        random.randint(300000, 5000000),
        random.choice(kyc_levels),
        pep_flag,
        aml_risk,
        open_date.date().isoformat(),
        (open_date - timedelta(days=random.randint(0, 3650))).date().isoformat()
    ])

    # Create 1-3 accounts per customer
    num_accounts = random.choices([1, 2, 3], weights=[60, 30, 10], k=1)[0]
    for _ in range(num_accounts):
        account_id = f"ACCT{account_counter:07d}"
        account_counter += 1
        acct_open = open_date + timedelta(days=random.randint(0, 1200))
        balance = random.randint(5000, 5000000)
        status = random.choice(["Active", "Dormant", "Closed"] if random.random() < 0.1 else ["Active"])
        accounts_ws.append([
            account_id,
            customer_id,
            random.choice(account_types),
            customer_currency,
            random.choice(customer_branches),
            balance,
            acct_open.date().isoformat(),
            status
        ])

        # Transactions for each account
        num_transactions = random.randint(20, 100)
        for _ in range(num_transactions):
            txn_date = acct_open + timedelta(days=random.randint(0, 2000))
            txn_type = random.choice(transaction_types)
            amount = random.randint(100, 2000000)
            if txn_type == "Debit" and amount > balance:
                amount = random.randint(100, min(balance, 200000))
            transaction_rows.append([
                f"TXN{transaction_counter:08d}",
                account_id,
                customer_id,
                txn_date.date().isoformat(),
                amount,
                customer_currency,
                txn_type,
                random.choice(["Vendor", "Employer", "Family", "Utility", "Unknown"]),
                random.choice(channels),
                random.choice(customer_branches)
            ])
            transaction_counter += 1

        # Alerts from suspicious accounts or customers
        if pep_flag or esg_flag or balance > 3000000 or random.random() < 0.03:
            alert_type = random.choice(["AML", "KYC", "Fraud", "Sanctions"])
            severity = random.choice(["Low", "Medium", "High"])
            alert_reason = []
            if pep_flag:
                alert_reason.append("PEP flagged")
            if esg_flag:
                alert_reason.append("ESG violation")
            if balance > 3000000:
                alert_reason.append("High balance")
            if random.random() < 0.5:
                alert_reason.append("Unusual transaction pattern")
            alert_rows.append([
                f"ALERT{alert_counter:06d}",
                customer_id,
                account_id,
                (acct_open + timedelta(days=random.randint(0, 2000))).date().isoformat(),
                alert_type,
                severity,
                "; ".join(alert_reason) or "Review required"
            ])
            alert_counter += 1

    kyc_rows.append([
        customer_id,
        (open_date + timedelta(days=random.randint(0, 365))).date().isoformat(),
        random.choice(["Verified", "Pending", "Review"]),
        random.choice(["ID, Address", "ID, Address, Income Proof", "ID"]),
        aml_risk,
        "Automated review" if aml_risk == "Low" else "Manual review"
    ])

# Append transaction and alert rows
for row in transaction_rows:
    transactions_ws.append(row)

for row in alert_rows:
    alerts_ws.append(row)

for row in kyc_rows:
    kyc_ws.append(row)

path = "/{folder_path}/ai-intelligence-network-team-delta/backend/synthetic_data/customer_dataset.xlsx"
wb.save(path)
print(path)